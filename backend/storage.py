from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook


BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
WORKBOOK_PATH = DATA_DIR / "expenses.xlsx"
SHEET_NAME = "expenses"
HEADERS = [
    "id",
    "date",
    "category",
    "description",
    "amount",
    "payment_method",
    "notes",
]


@dataclass
class Expense:
    id: str
    date: str
    category: str
    description: str
    amount: float
    payment_method: str
    notes: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "date": self.date,
            "category": self.category,
            "description": self.description,
            "amount": round(float(self.amount), 2),
            "payment_method": self.payment_method,
            "notes": self.notes,
        }


def ensure_workbook() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if WORKBOOK_PATH.exists():
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    workbook.save(WORKBOOK_PATH)


def _load_sheet():
    ensure_workbook()
    workbook = load_workbook(WORKBOOK_PATH)
    sheet = workbook[SHEET_NAME]
    return workbook, sheet


def list_expenses() -> list[dict[str, Any]]:
    _, sheet = _load_sheet()
    expenses: list[Expense] = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        expense = Expense(
            id=str(row[0] or ""),
            date=_serialize_date(row[1]),
            category=str(row[2] or ""),
            description=str(row[3] or ""),
            amount=float(row[4] or 0),
            payment_method=str(row[5] or ""),
            notes=str(row[6] or ""),
        )
        expenses.append(expense)

    expenses.sort(key=lambda item: (item.date, item.id), reverse=True)
    return [expense.as_dict() for expense in expenses]


def add_expense(payload: dict[str, Any]) -> dict[str, Any]:
    workbook, sheet = _load_sheet()
    expense = Expense(
        id=str(uuid4()),
        date=_normalize_date(payload.get("date")),
        category=_require_text(payload.get("category"), "category"),
        description=_require_text(payload.get("description"), "description"),
        amount=_normalize_amount(payload.get("amount")),
        payment_method=_require_text(payload.get("payment_method"), "payment_method"),
        notes=str(payload.get("notes", "")).strip(),
    )
    sheet.append(
        [
            expense.id,
            expense.date,
            expense.category,
            expense.description,
            expense.amount,
            expense.payment_method,
            expense.notes,
        ]
    )
    workbook.save(WORKBOOK_PATH)
    return expense.as_dict()


def delete_expense(expense_id: str) -> bool:
    workbook, sheet = _load_sheet()
    for row_index in range(2, sheet.max_row + 1):
        if str(sheet.cell(row=row_index, column=1).value or "") == expense_id:
            sheet.delete_rows(row_index, 1)
            workbook.save(WORKBOOK_PATH)
            return True
    return False


def calculate_stats(expenses: list[dict[str, Any]]) -> dict[str, Any]:
    total_spend = round(sum(item["amount"] for item in expenses), 2)
    average_transaction = round(total_spend / len(expenses), 2) if expenses else 0.0

    today = date.today()
    current_month_key = today.strftime("%Y-%m")
    current_month_expenses = [item for item in expenses if item["date"].startswith(current_month_key)]
    current_month_total = round(sum(item["amount"] for item in current_month_expenses), 2)

    category_totals: dict[str, float] = {}
    month_totals: dict[str, float] = {}

    for item in expenses:
        category_totals[item["category"]] = category_totals.get(item["category"], 0.0) + item["amount"]
        month_key = item["date"][:7]
        month_totals[month_key] = month_totals.get(month_key, 0.0) + item["amount"]

    top_category = ""
    if category_totals:
        top_category = max(category_totals.items(), key=lambda pair: pair[1])[0]

    monthly_series = [
        {"label": month, "total": round(total, 2)}
        for month, total in sorted(month_totals.items())[-6:]
    ]

    category_series = [
        {"label": category, "total": round(total, 2)}
        for category, total in sorted(category_totals.items(), key=lambda pair: pair[1], reverse=True)
    ]

    return {
        "cards": {
            "total_spend": total_spend,
            "current_month_total": current_month_total,
            "average_transaction": average_transaction,
            "top_category": top_category or "N/A",
            "transaction_count": len(expenses),
        },
        "monthly_series": monthly_series,
        "category_series": category_series,
    }


def _serialize_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "")


def _normalize_date(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        raise ValueError("date is required")
    return datetime.fromisoformat(raw).date().isoformat()


def _normalize_amount(value: Any) -> float:
    try:
        amount = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("amount must be numeric") from exc

    if amount <= 0:
        raise ValueError("amount must be greater than zero")
    return round(amount, 2)


def _require_text(value: Any, field_name: str) -> str:
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"{field_name} is required")
    return text
